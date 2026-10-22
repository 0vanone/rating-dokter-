import io

from django.utils import timezone

from . import services
from .cli_utils import aman_sel
from .models import Survei

HEADER_SURVEI = ["ID", "Waktu", "Poliklinik", "Dokter", "Nama pasien", "No HP",
                 "Tanggal kunjungan", "Skor rata-rata", "Saran"]


def _skor(nilai):
    return None if nilai is None else round(float(nilai), 2)


def _label_periode(kode):
    return dict(services.PERIODE).get(kode, "Semua waktu")


def buat_xlsx(kode_periode="semua"):
    """Laporan Excel: ringkasan, peringkat dokter, dan seluruh data survei pada periode itu."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    sejak = services.mulai_periode(kode_periode)
    ringkas = services.ringkasan(sejak)
    tebal = Font(bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["Laporan Kepuasan Pasien"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Periode", _label_periode(kode_periode)])
    ws.append(["Total survei masuk", ringkas["total"]])
    ws.append(["Skor rerata RS (1.0-5.0)", _skor(ringkas["rata"])])
    ws.append(["Komplain aktif", ringkas["komplain_aktif"]])
    ws.append([])
    ws.append(["Poliklinik", "Responden", "Skor rerata"])
    for sel in ws[ws.max_row]:
        sel.font = tebal
    for p in services.per_poli(sejak):
        ws.append([p.nama, p.jumlah, _skor(p.rata)])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14

    ws2 = wb.create_sheet("Peringkat Dokter")
    ws2.append(["#", "Dokter", "Poliklinik", "Responden", "Skor rerata"])
    for sel in ws2[1]:
        sel.font = tebal
    for i, d in enumerate(services.leaderboard(sejak=sejak), start=1):
        ws2.append([i, d.nama, d.poliklinik.nama, d.jumlah, _skor(d.rata)])
    for kolom, lebar in zip("ABCDE", (5, 32, 20, 12, 12)):
        ws2.column_dimensions[kolom].width = lebar

    ws3 = wb.create_sheet("Data Survei")
    ws3.append(HEADER_SURVEI)
    for sel in ws3[1]:
        sel.font = tebal
    qs = Survei.objects.select_related("dokter__poliklinik").order_by("dibuat")
    if sejak:
        qs = qs.filter(dibuat__gte=sejak)
    for s in qs:
        ws3.append([
            s.pk,
            timezone.localtime(s.dibuat).strftime("%Y-%m-%d %H:%M"),
            s.dokter.poliklinik.nama,
            s.dokter.nama,
            aman_sel(s.nama_pasien),
            aman_sel(s.no_hp),
            s.tanggal_kunjungan.isoformat(),
            _skor(s.skor_rata),
            aman_sel(s.saran),
        ])
    for kolom, lebar in zip("ABCDEFGHI", (7, 18, 18, 28, 24, 16, 18, 14, 60)):
        ws3.column_dimensions[kolom].width = lebar

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def buat_pdf(kode_periode="semua"):
    """Laporan PDF ringkas untuk bahan rapat direksi."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    sejak = services.mulai_periode(kode_periode)
    ringkas = services.ringkasan(sejak)
    gaya = getSampleStyleSheet()

    def tabel(data, lebar):
        t = Table(data, colWidths=lebar, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b7a75")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f7f8")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cfdcdf")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        return t

    isi = [
        Paragraph("Laporan Kepuasan Pasien", gaya["Title"]),
        Paragraph(f"Periode: {_label_periode(kode_periode)}. Dibuat {timezone.localtime():%d-%m-%Y %H:%M}.", gaya["Normal"]),
        Spacer(1, 0.5 * cm),
        tabel(
            [["Total survei masuk", "Skor rerata RS (1.0-5.0)", "Komplain aktif"],
             [str(ringkas["total"]), "-" if ringkas["rata"] is None else f"{float(ringkas['rata']):.2f}",
              str(ringkas["komplain_aktif"])]],
            [5.5 * cm, 5.5 * cm, 5.5 * cm],
        ),
        Spacer(1, 0.6 * cm),
        Paragraph("Skor per poliklinik", gaya["Heading3"]),
        tabel(
            [["Poliklinik", "Responden", "Skor rerata"]]
            + [[p.nama, str(p.jumlah), "-" if p.rata is None else f"{float(p.rata):.2f}"]
               for p in services.per_poli(sejak)],
            [8 * cm, 4 * cm, 4.5 * cm],
        ),
        Spacer(1, 0.6 * cm),
        Paragraph("Peringkat dokter", gaya["Heading3"]),
    ]
    baris = [[str(i), d.nama, d.poliklinik.nama, str(d.jumlah), f"{float(d.rata):.2f}"]
             for i, d in enumerate(services.leaderboard(sejak=sejak), start=1)]
    if baris:
        isi.append(tabel([["#", "Dokter", "Poliklinik", "Responden", "Skor"]] + baris,
                         [1 * cm, 6.5 * cm, 4.5 * cm, 2.5 * cm, 2 * cm]))
    else:
        isi.append(Paragraph("Belum ada data survei pada periode ini.", gaya["Normal"]))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm, title="Laporan Kepuasan Pasien").build(isi)
    return buf.getvalue()


def _qr_image(url):
    import qrcode
    import qrcode.constants

    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_H, box_size=20, border=4)
    qr.add_data(url)
    qr.make(fit=True)
    return qr.make_image(fill_color="black", back_color="white").get_image().convert("RGB")


def buat_qr_png(url):
    buf = io.BytesIO()
    _qr_image(url).save(buf, format="PNG")
    return buf.getvalue()


def buat_qr_pdf(url):
    """PDF satu halaman berisi QR, siap cetak untuk papan nama, meja pendaftaran, atau banner."""
    buf = io.BytesIO()
    _qr_image(url).save(buf, format="PDF", resolution=150.0)
    return buf.getvalue()
